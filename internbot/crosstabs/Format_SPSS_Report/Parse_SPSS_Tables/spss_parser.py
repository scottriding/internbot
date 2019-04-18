from openpyxl import load_workbook, Workbook
from collections import OrderedDict

class SPSSParser(object):

    def __init__(self):
        self.__tables = []
        self.__responses = []
        self.__base_size_row = 1
        self.__freq_table_row = 1
        self.__sign_table_row = 1
        self.__significant_defintion = []

        # calculate the excel alphabet from D to ZZZ
        alphabet = []
        for letter in range(68, 91):
            alphabet.append(chr(letter))

        self.__extend_alphabet = []
        self.__extend_alphabet.extend(alphabet)
        full_alphabet = []
        double_alphabet = []
        for letter in range(65, 91):
            full_alphabet.append(chr(letter))
        index = 0
        while index < len(full_alphabet):
            for letter in full_alphabet:
                double_letters = "%s%s" % (full_alphabet[index], letter)
                double_alphabet.append(double_letters)
                self.__extend_alphabet.append(double_letters)
            index += 1
        index = 0
        while index < len(double_alphabet):
            for letter in full_alphabet:
                triple_letters = "%s%s" % (double_alphabet[index], letter)
                self.__extend_alphabet.append(triple_letters)
            index += 1

    def get_tables(self):
        return self.__tables

    def add_table(self, path_to_workbook):
        self.__responses = []
        self.__base_size_row = 1
        self.__freq_table_row = 1
        self.__sign_table_row = 1
        self.__banners = []
        self.__significant_defintion = []

        print "Loading next workbook"
        current_workbook = load_workbook(path_to_workbook)
        current_sheet = current_workbook.get_active_sheet()

        table_name = current_sheet["A1"].value
        print "Parsing: %s" % (table_name)
        table_base_desc = current_sheet["A2"].value

        self.parse_table(current_sheet)

        base_size_cell = "C%s" % self.__base_size_row
        base_size = current_sheet[base_size_cell].value
        
        new_table = Table(table_name, table_base_desc, base_size, self.__banners, self.__responses, self.total_row, self.__significant_defintion)
        self.__tables.append(new_table)

    def parse_table(self, sheet):
        self.configure_table_rows(sheet)
        self.parse_banners(sheet)
        self.parse_responses(sheet)
        self.parse_banner_pts(sheet)
        self.parse_significance_table(sheet)
        self.parse_total_row(sheet)
        self.parse_sig_cells(sheet)
        self.parse_sig_definition(sheet)

    def configure_table_rows(self, sheet):
        current_row = 1
        top_of_table = True
        for cell in sheet["B"]:
            if cell.value != 'Sigma' and cell.value is not None:
                if top_of_table is True:
                    self.__freq_table_row = current_row
                    top_of_table = False

            if cell.value == 'Sigma':
                self.__base_size_row = current_row
                break

            current_row += 1

        self.__sign_table_row = self.__base_size_row + 2

    def parse_banners(self, sheet):

        self.__banner_depth = self.__freq_table_row - 2

        # faux switch case - cover up to a depth of 8
        if(self.__banner_depth == 8):
            self.banner_parse_eight(sheet)
        elif(self.__banner_depth == 7):
            self.banner_parse_seven(sheet)
        elif(self.__banner_depth == 6):
            self.banner_parse_six(sheet)
        elif(self.__banner_depth == 5):
            self.banner_parse_five(sheet)
        elif(self.__banner_depth == 4):
            self.banner_parse_four(sheet)
        elif(self.__banner_depth == 3):
            self.banner_parse_three(sheet)
        elif(self.__banner_depth == 2):
            self.banner_parse_two(sheet)
        elif(self.__banner_depth == 1):
            self.banner_parse_one(sheet)
        else:
            raise Exception("Could not parse banners")

    def banner_parse_eight(self, sheet):
        all_banners = []
        current_high_level = ""
        current_sub_level_one = ""
        current_sub_level_two = ""
        current_sub_level_three = ""
        current_sub_level_four = ""
        current_sub_level_five = ""
        current_sub_level_six = ""
        for letter in self.__extend_alphabet:
            banner_row = 2
            current_banner_pt = []

            # topline level banner
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_high_level = sheet[banner_cell].value
            current_banner_pt.append(current_high_level)

            # level 1
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_one = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_one)

            # level 2
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_two = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_two)

            # level 3
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_three = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_three)

            # level 4
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_four = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_four)

            # level 5
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_five = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_five)

            # level 6
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_six = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_six)

            # banner point level
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_banner_pt.append(sheet[banner_cell].value)
            else:
                break

            all_banners.append(current_banner_pt)
        self.__banners = all_banners

    def banner_parse_seven(self, sheet):
        all_banners = []
        current_high_level = ""
        current_sub_level_one = ""
        current_sub_level_two = ""
        current_sub_level_three = ""
        current_sub_level_four = ""
        current_sub_level_five = ""
        for letter in self.__extend_alphabet:
            banner_row = 2
            current_banner_pt = []

            # topline level banner
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_high_level = sheet[banner_cell].value
            current_banner_pt.append(current_high_level)

            # level 1
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_one = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_one)

            # level 2
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_two = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_two)

            # level 3
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_three = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_three)

            # level 4
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_four = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_four)

            # level 5
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_five = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_five)

            # banner point level
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_banner_pt.append(sheet[banner_cell].value)
            else:
                break

            all_banners.append(current_banner_pt)
        self.__banners = all_banners

    def banner_parse_six(self, sheet):
        all_banners = []
        current_high_level = ""
        current_sub_level_one = ""
        current_sub_level_two = ""
        current_sub_level_three = ""
        current_sub_level_four = ""
        for letter in self.__extend_alphabet:
            banner_row = 2
            current_banner_pt = []

            # topline level banner
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_high_level = sheet[banner_cell].value
            current_banner_pt.append(current_high_level)

            # level 1
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_one = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_one)

            # level 2
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_two = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_two)

            # level 3
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_three = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_three)

            # level 4
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_four = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_four)

            # banner point level
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_banner_pt.append(sheet[banner_cell].value)
            else:
                break

            all_banners.append(current_banner_pt)
        self.__banners = all_banners

    def banner_parse_five(self, sheet):
        all_banners = []
        current_high_level = ""
        current_sub_level_one = ""
        current_sub_level_two = ""
        current_sub_level_three = ""
        for letter in self.__extend_alphabet:
            banner_row = 2
            current_banner_pt = []

            # topline level banner
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_high_level = sheet[banner_cell].value
            current_banner_pt.append(current_high_level)

            # level 1
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_one = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_one)

            # level 2
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_two = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_two)

            # level 3
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_three = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_three)

            # banner point level
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_banner_pt.append(sheet[banner_cell].value)
            else:
                break

            all_banners.append(current_banner_pt)
        self.__banners = all_banners

    def banner_parse_four(self, sheet):
        all_banners = []
        current_high_level = ""
        current_sub_level_one = ""
        current_sub_level_two = ""
        for letter in self.__extend_alphabet:
            banner_row = 2
            current_banner_pt = []

            # topline level banner
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_high_level = sheet[banner_cell].value
            current_banner_pt.append(current_high_level)

            # level 1
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_one = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_one)

            # level 2
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_two = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_two)

            # banner point level
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_banner_pt.append(sheet[banner_cell].value)
            else:
                break

            all_banners.append(current_banner_pt)
        self.__banners = all_banners

    def banner_parse_three(self, sheet):
        all_banners = []
        current_high_level = ""
        current_sub_level_one = ""
        for letter in self.__extend_alphabet:
            banner_row = 2
            current_banner_pt = []

            # topline level banner
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_high_level = sheet[banner_cell].value
            current_banner_pt.append(current_high_level)

            # level 1
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_sub_level_one = sheet[banner_cell].value
            current_banner_pt.append(current_sub_level_one)

            # banner point level
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_banner_pt.append(sheet[banner_cell].value)
            else:
                break

            all_banners.append(current_banner_pt)
        self.__banners = all_banners

    def banner_parse_two(self, sheet):
        all_banners = []
        current_high_level = ""
        current_sub_level_one = ""
        for letter in self.__extend_alphabet:
            banner_row = 2
            current_banner_pt = []

            # topline level banner
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_high_level = sheet[banner_cell].value
            current_banner_pt.append(current_high_level)

            # banner point level
            banner_row += 1
            banner_details = []
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_banner_pt.append(sheet[banner_cell].value)
            else:
                break

            all_banners.append(current_banner_pt)
        self.__banners = all_banners

    def banner_parse_one(self, sheet):
        all_banners = []
        current_high_level = ""
        current_sub_level_one = ""
        for letter in self.__extend_alphabet:
            banner_row = 2
            current_banner_pt = []

            # banner point level
            banner_cell = "%s%s" % (letter, str(banner_row))
            if sheet[banner_cell].value is not None:
                current_banner_pt.append(sheet[banner_cell].value)
            else:
                break
            all_banners.append(current_banner_pt)
        self.__banners = all_banners

    def parse_responses(self, sheet):
        current_row = self.__freq_table_row
        while current_row < self.__base_size_row:
            current_cell = "B%s" % current_row
            
            if sheet[current_cell].value is not None:
                response_name = sheet[current_cell].value

                response_pop_cell = "C%s" % current_row
                response_perc_cell = "C%s" % str(current_row + 1)

                response_pop = sheet[response_pop_cell].value
                response_perc = sheet[response_perc_cell].value
                self.__responses.append(Response(response_name, response_pop, response_perc, current_row))

            current_row += 1

    def parse_banner_pts(self, sheet):
        for letter in self.__extend_alphabet:
            banner_pt_row = self.__freq_table_row - 1
            banner_pt_cell = "%s%s" % (letter, banner_pt_row)
            if sheet[banner_pt_cell].value is None:
                break
            else:
                banner_pt_name = sheet[banner_pt_cell].value
                for response in self.__responses:
                    current_perc_cell = "%s%s" % (letter, str(response.row + 1))
                    current_pop_cell = "%s%s" % (letter, str(response.row))
            
                    percentage = sheet[current_perc_cell].value
                    population = sheet[current_pop_cell].value

                    response.add_banner_pt(banner_pt_name, population, percentage)

    def parse_significance_table(self, sheet):
        current_row = self.__sign_table_row
        assign_sig_banner = True
        index = 0
        while True:
            current_cell = "B%s" % current_row
            if sheet[current_cell].value is None and index >= len(self.__responses):
                break
            elif sheet[current_cell].value is None:
                current_row += 1
            elif sheet[current_cell].value is not None:
                if assign_sig_banner is True:
                    self.__sign_table_row = current_row - 2
                    assign_sig_banner = False

                current_response = self.__responses[index]

                significant_cell = "C%s" % current_row
                current_response.sig_details = sheet[significant_cell].value

                current_row += 1
                index += 1

    def parse_total_row(self, sheet):
        number_of_cols = len(self.__responses[0].banner_pts)
        self.total_row = TotalRow()
        current_row = 1
        for cell in sheet["B"]:
            if cell.value is not None and cell.value == "Sigma":
                break
            current_row += 1

        letter_index = 0
        while letter_index < number_of_cols:
            population_cell = "%s%s" % (self.__extend_alphabet[letter_index], current_row)
            percentage_cell = "%s%s" % (self.__extend_alphabet[letter_index], current_row + 1)
            self.total_row.add_population(sheet[population_cell].value)
            self.total_row.add_percentage(sheet[percentage_cell].value)
            letter_index += 1

    def parse_sig_cells(self, sheet):
        banner_index = 0
        for letter in self.__extend_alphabet:
            banner_pt_row = self.__sign_table_row
            banner_pt_cell = "%s%s" % (letter, banner_pt_row)
            if sheet[banner_pt_cell].value is None:
                break
            else:
                response_row = self.__sign_table_row + 2
                for response in self.__responses:
                    current_cell = "%s%s" % (letter, str(response_row))
                    significant_value = sheet[current_cell].value
                    response.banner_pts[banner_index].sig_details = significant_value
                    response_row += 1
                banner_index += 1
                self.__sig_def = response_row
  
    def parse_sig_definition(self, sheet):
        current_row = self.__sig_def
        current_cell = "A%s" % current_row
        while sheet[current_cell].value is not None:
            self.__significant_defintion.append(sheet[current_cell].value)
            current_row += 1
            current_cell = "A%s" % current_row

class Table(object):

    def __init__(self, name, base, base_size, banners, responses, total_row, sig_desc):
        self.__name = name
        self.__base = base
        self.__base_size = base_size
        self.__banners = banners
        self.__responses = responses
        self.__total_row = total_row
        self.__sig_desc = sig_desc

    @property
    def name(self):
        return self.__name

    @property
    def base_description(self):
        return self.__base

    @property
    def base_size(self):
        return self.__base_size

    @property
    def banners(self):
        return self.__banners

    @property
    def responses(self):
        return self.__responses

    @property
    def total_row(self):
        return self.__total_row

    @property
    def sig_desc(self):
        return self.__sig_desc

    @property
    def count_banner_pts(self):
        return len(self.__responses[0].banner_pts)
                  
class Response(object):

    def __init__(self, name, pop, perc, row):
        self.__name = name
        self.__population = pop
        self.__percentage = perc
        self.__row = row
        self.__sig_details = ""
        self.__banner_pts = []

    @property
    def name(self):
        return self.__name

    @property
    def row(self):
        return self.__row

    @property
    def population(self):
        return self.__population

    @property
    def percentage(self):
        return self.__percentage

    @property
    def banner_pts(self):
        return self.__banner_pts

    @property
    def sig_details(self):
        return self.__sig_details

    @sig_details.setter
    def sig_details(self, details):
        self.__sig_details = details

    def add_banner_pt(self, banner_name, banner_pop, banner_perc):
        self.__banner_pts.append(BannerPt(banner_name, banner_pop, banner_perc))

    def __repr__ (self):
        result = ''
        result += "Response: %s \n" % (self.__name)
        result += "Significance: %s \n" % self.__sig_details
        return result

class BannerPt(object):

    def __init__(self, name, pop, perc):
        self.__name = name
        self.__population = pop
        self.__percentage = perc
        self.__sig_details = ""

    @property
    def name(self):
        return self.__name

    @property
    def sig_details(self):
        return self.__sig_details

    @property
    def population(self):
        return self.__population

    @property
    def percentage(self):
        return self.__percentage

    @sig_details.setter
    def sig_details(self, details):
        self.__sig_details = details

    def __repr__(self):
        result = "%s\n" % self.__name
        return result

class TotalRow(object):

    def __init__(self):
        self.__populations = []
        self.__percentages = []

    def add_population(self, population):
        self.__populations.append(population)

    def add_percentage(self, percentage):
        self.__percentages.append(percentage)

    @property
    def populations(self):
        return self.__populations

    @property
    def percentages(self):
        return self.__percentages
